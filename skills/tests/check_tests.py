# -*- coding: utf-8 -*-
"""Check what a skill actually did to a test file.

Usage:
    python check_tests.py <dir>            # capture baselines, before running skills
    python check_tests.py <dir> --verify   # compare against the baselines

Reports PASS/FAIL per rule. The traps are the point: a skill that trims the
overloaded slides but also mangles the table slide has failed, even though
the happy path looks perfect.
"""
import json
import os
import re
import sys
import zipfile

sys.stdout.reconfigure(encoding="utf-8")

import openpyxl

HEB = re.compile(r"[֐-׿]")


def slide_texts(path):
    z = zipfile.ZipFile(path)
    names = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                   key=lambda n: int(re.search(r"\d+", n.split("/")[-1]).group()))
    out = []
    for n in names:
        xml = z.read(n).decode("utf-8")
        out.append({
            "lines": [t for t in re.findall(r"<a:t>(.*?)</a:t>", xml, re.S) if t.strip()],
            "has_table": "<a:tbl>" in xml,
            "has_chart": "graphicFrame" in xml and "chart" in xml,
        })
    notes = {}
    for n in z.namelist():
        m = re.match(r"ppt/notesSlides/notesSlide(\d+)\.xml$", n)
        if m:
            xml = z.read(n).decode("utf-8")
            notes[int(m.group(1))] = " ".join(
                t for t in re.findall(r"<a:t>(.*?)</a:t>", xml, re.S) if t.strip())
    return {"slides": out, "notes": notes}


def xlsx_shape(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    return {"sheet": ws.title, "max_col": ws.max_column, "max_row": ws.max_row,
            "header": rows[0] if rows else [], "rows": rows}


def capture(d):
    base = {}
    for f in ("test-slide-diet.pptx", "test-speaker-notes.pptx"):
        p = os.path.join(d, f)
        if os.path.exists(p):
            base[f] = slide_texts(p)
    for f in ("test-variance.xlsx", "test-variance-malformed.xlsx"):
        p = os.path.join(d, f)
        if os.path.exists(p):
            base[f] = xlsx_shape(p)
    dest = os.path.join(d, "baseline.json")
    json.dump(base, open(dest, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("baseline saved ->", dest)
    for f, v in base.items():
        if "slides" in v:
            print(f"  {f}: {len(v['slides'])} slides, "
                  f"{sum(len(s['lines']) for s in v['slides'])} text runs, "
                  f"{len(v['notes'])} notes pages")
        else:
            print(f"  {f}: {v['max_row']} rows x {v['max_col']} cols")
    return 0


def rule(ok, label, detail=""):
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}" + (f" — {detail}" if detail else ""))
    return 0 if ok else 1


def verify(d):
    base = json.load(open(os.path.join(d, "baseline.json"), encoding="utf-8"))
    fails = 0

    # ---------- slide-diet ----------
    f = "test-slide-diet.pptx"
    p = os.path.join(d, f)
    if f in base and os.path.exists(p):
        print(f"\n=== {f} ===")
        b, a = base[f], slide_texts(p)
        # trap 1: table slide untouched
        fails += rule(b["slides"][1]["lines"] == a["slides"][1]["lines"],
                      "slide 2 (table) untouched")
        # trap 2: chart slide untouched
        fails += rule(b["slides"][2]["lines"] == a["slides"][2]["lines"],
                      "slide 3 (chart) untouched")
        # trap 3: already-compliant slide untouched
        fails += rule(b["slides"][3]["lines"] == a["slides"][3]["lines"],
                      "slide 4 (already compliant) untouched")
        # slide 1 should have been trimmed
        b1 = len(" ".join(b["slides"][0]["lines"]).split())
        a1 = len(" ".join(a["slides"][0]["lines"]).split())
        fails += rule(a1 < b1, "slide 1 trimmed", f"{b1} -> {a1} words on slide")
        # facts must survive somewhere in the file
        facts = ["535", "567", "101", "103", "104", "שירה ברק", "28.09.2026", "285"]
        alltext = " ".join(t for s in a["slides"] for t in s["lines"]) + " " + " ".join(a["notes"].values())
        missing = [x for x in facts if x not in alltext]
        fails += rule(not missing, "no fact lost (slide or notes)",
                      f"missing: {missing}" if missing else "all 8 present")

    # ---------- speaker-notes-he ----------
    f = "test-speaker-notes.pptx"
    p = os.path.join(d, f)
    if f in base and os.path.exists(p):
        print(f"\n=== {f} ===")
        b, a = base[f], slide_texts(p)
        alltext = " ".join(a["notes"].values())
        lost = [f"MARKER-{i:02d}" for i in range(1, 5) if f"MARKER-{i:02d}" not in alltext]
        fails += rule(not lost, "existing notes preserved",
                      f"lost: {lost}" if lost else "all 4 markers survive")
        vals = [str(i * 1111) for i in range(1, 5)]
        lostv = [v for v in vals if v not in alltext]
        fails += rule(not lostv, "unique values in old notes survive",
                      f"lost: {lostv}" if lostv else "all 4 present")
        grew = sum(len(v) for v in a["notes"].values()) > sum(len(v) for v in b["notes"].values())
        fails += rule(grew, "notes were actually written")
        for i, s in enumerate(a["slides"]):
            if s["lines"] != b["slides"][i]["lines"]:
                fails += rule(False, f"slide {i+1} body unchanged", "slide text was modified")
                break
        else:
            fails += rule(True, "no slide body was modified")

    # ---------- variance-analysis ----------
    f = "test-variance.xlsx"
    p = os.path.join(d, f)
    if f in base and os.path.exists(p):
        print(f"\n=== {f} ===")
        b, a = base[f], xlsx_shape(p)
        fails += rule(a["max_col"] > b["max_col"], "columns added",
                      f"{b['max_col']} -> {a['max_col']}")
        fails += rule(a["max_row"] >= b["max_row"], "no rows lost",
                      f"{b['max_row']} -> {a['max_row']}")
        flat = " ".join(str(c) for r in a["rows"] for c in r if c is not None)
        fails += rule("#DIV/0!" not in flat, "no #DIV/0! in the sheet")
        for i in range(min(len(b["rows"]), len(a["rows"]))):
            orig = b["rows"][i]
            now = a["rows"][i][:len(orig)]
            if orig != now:
                fails += rule(False, "original columns intact",
                              f"row {i+1} changed: {orig} -> {now}")
                break
        else:
            fails += rule(True, "original columns intact")

    print("\n" + ("ALL RULES PASSED" if fails == 0 else f"{fails} RULE(S) FAILED"))
    return 1 if fails else 0


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    d = sys.argv[1]
    return verify(d) if "--verify" in sys.argv else capture(d)


if __name__ == "__main__":
    sys.exit(main())
